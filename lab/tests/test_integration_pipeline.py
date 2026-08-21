"""Integration and E2E tests for the full Mythos pipeline.

Validates cross-module flows:
- Parser → ReviewEngine → Export (full review pipeline)
- Parser → Reports → Export (rollover pipeline)
- Service.full_review() end-to-end
- Service.review() → export_review() chain
- 8858 full pipeline
- Filtered review pipeline
"""

import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from lab.xml_parser.api import MythosService
from lab.xml_parser.api.service import ReviewResult, ExportResult
from lab.xml_parser.core.filters import FilterSpec, build_filter_spec, filter_review_report
from lab.xml_parser.core.models import ReviewReport, RolloverReport
from lab.xml_parser.export import export_excel, export_csv, export_findings_csv
from lab.xml_parser.parser import EFileParser
from lab.xml_parser.reports import run_all_reports, run_all_reports_8858
from lab.xml_parser.review_engine import ReviewEngine

FIXTURES = Path(__file__).parent / "fixtures"


# ═══════════════════════════════════════════════════════════════════════════════
# E2E — Full Review Pipeline (XML → Checks → Export)
# ═══════════════════════════════════════════════════════════════════════════════


class TestFullReviewPipeline:
    """End-to-end: load XML, run all checks, export, verify output."""

    def test_review_and_export_excel(self, tmp_path):
        engine = ReviewEngine()
        report = engine.review(FIXTURES / "sample_cy.xml", FIXTURES / "sample_py.xml")
        out = tmp_path / "e2e_review.xlsx"
        export_excel(out, review_report=report, client_name="E2E Test")
        assert out.exists()
        assert out.stat().st_size > 1000

    def test_review_and_export_csv(self, tmp_path):
        engine = ReviewEngine()
        report = engine.review(FIXTURES / "sample_cy.xml", FIXTURES / "sample_py.xml")
        out = tmp_path / "e2e_findings.csv"
        export_findings_csv(report, out)
        df = pd.read_csv(out)
        assert len(df) == 28
        assert set(df["severity"].unique()) == {"HIGH", "MEDIUM", "LOW"}

    def test_full_review_via_service(self, tmp_path):
        svc = MythosService()
        reports = svc.full_review(
            FIXTURES / "sample_cy.xml",
            FIXTURES / "sample_py.xml",
            output_path=tmp_path / "full_review.xlsx",
        )
        assert isinstance(reports, dict)
        assert "_review_report" in reports
        assert "_export_path" in reports
        assert reports["_export_path"].exists()

        review_report = reports["_review_report"]
        assert len(review_report.findings) == 28

    def test_full_review_rollover_reports_complete(self, tmp_path):
        svc = MythosService()
        reports = svc.full_review(
            FIXTURES / "sample_cy.xml",
            FIXTURES / "sample_py.xml",
            output_path=tmp_path / "full.xlsx",
        )
        report_keys = [k for k in reports if not k.startswith("_")]
        assert len(report_keys) == 12
        for key in report_keys:
            assert isinstance(reports[key], RolloverReport)


# ═══════════════════════════════════════════════════════════════════════════════
# E2E — Rollover Reports Pipeline
# ═══════════════════════════════════════════════════════════════════════════════


class TestRolloverPipeline:
    """End-to-end: load XML pair, generate all reports, export, verify."""

    def test_rollover_reports_to_csv(self, tmp_path):
        reports = run_all_reports(FIXTURES / "sample_py.xml", FIXTURES / "sample_cy.xml")
        files = export_csv(reports, tmp_path, prefix="e2e")
        assert len(files) >= 9
        total_rows = 0
        for f in files:
            df = pd.read_csv(f)
            total_rows += len(df)
        assert total_rows > 30

    def test_rollover_reports_to_excel(self, tmp_path):
        reports = run_all_reports(FIXTURES / "sample_py.xml", FIXTURES / "sample_cy.xml")
        out = tmp_path / "rollover.xlsx"
        export_excel(out, rollover_reports=reports, client_name="E2E Test")
        assert out.exists()
        assert out.stat().st_size > 5000

    def test_rollover_pass_fail_totals_consistent(self):
        reports = run_all_reports(FIXTURES / "sample_py.xml", FIXTURES / "sample_cy.xml")
        for name, report in reports.items():
            assert report.passed + report.failed == report.total_checks, (
                f"{name}: passed({report.passed}) + failed({report.failed}) != total({report.total_checks})"
            )


# ═══════════════════════════════════════════════════════════════════════════════
# E2E — Form 8858 Pipeline
# ═══════════════════════════════════════════════════════════════════════════════


class TestPipeline8858:
    """End-to-end for Form 8858 flow."""

    def test_8858_review_and_export(self, tmp_path):
        engine = ReviewEngine()
        report = engine.review(
            FIXTURES / "sample_8858_cy.xml",
            FIXTURES / "sample_8858_py.xml",
            form_type="8858",
        )
        assert report.form_type == "8858"
        assert len(report.findings) >= 5

        out = tmp_path / "8858_findings.csv"
        export_findings_csv(report, out)
        df = pd.read_csv(out)
        assert len(df) == len(report.findings)

    def test_8858_reports_pipeline(self, tmp_path):
        reports = run_all_reports_8858(
            FIXTURES / "sample_8858_py.xml",
            FIXTURES / "sample_8858_cy.xml",
        )
        assert len(reports) == 5
        files = export_csv(reports, tmp_path, prefix="8858")
        assert len(files) == 5


# ═══════════════════════════════════════════════════════════════════════════════
# INTEGRATION — Service Layer
# ═══════════════════════════════════════════════════════════════════════════════


class TestServiceIntegration:
    """MythosService as entry point for full workflows."""

    def test_review_then_export_excel(self, tmp_path):
        svc = MythosService()
        result = svc.review(FIXTURES / "sample_cy.xml", prior=FIXTURES / "sample_py.xml")
        assert result.success
        export = svc.export_review(result.report, path=tmp_path / "svc.xlsx")
        assert export.success
        assert export.path.exists()

        import openpyxl
        wb = openpyxl.load_workbook(export.path, read_only=True)
        assert "Executive Summary" in wb.sheetnames
        assert "All Findings" in wb.sheetnames

    def test_review_then_export_csv(self, tmp_path):
        svc = MythosService()
        result = svc.review(FIXTURES / "sample_cy.xml", prior=FIXTURES / "sample_py.xml")
        export = svc.export_review(result.report, path=tmp_path / "svc.csv", format="csv")
        assert export.success
        df = pd.read_csv(export.path)
        assert len(df) == 28

    def test_review_then_export_json(self, tmp_path):
        svc = MythosService()
        result = svc.review(FIXTURES / "sample_cy.xml", prior=FIXTURES / "sample_py.xml")
        export = svc.export_review(result.report, path=tmp_path / "svc.json", format="json")
        assert export.success
        import json
        data = json.loads(export.path.read_text())
        assert len(data) == 28

    def test_service_progress_callback(self):
        events = []
        svc = MythosService(progress=lambda msg, pct: events.append((msg, pct)))
        svc.review(FIXTURES / "sample_cy.xml", prior=FIXTURES / "sample_py.xml")
        assert len(events) >= 2
        assert events[0][1] < events[-1][1]
        assert events[-1][1] == 1.0

    def test_service_error_handling(self):
        svc = MythosService()
        result = svc.review("nonexistent.xml")
        assert result.success is False
        assert result.message != ""


# ═══════════════════════════════════════════════════════════════════════════════
# INTEGRATION — Filtered Pipeline
# ═══════════════════════════════════════════════════════════════════════════════


class TestFilteredPipeline:
    """End-to-end with filters applied."""

    def test_filtered_review_by_entity(self, tmp_path):
        svc = MythosService()
        result = svc.review(FIXTURES / "sample_cy.xml", prior=FIXTURES / "sample_py.xml")
        spec = FilterSpec(entity_codes=frozenset({"E001"}))
        filtered = filter_review_report(result.report, spec)
        assert all(f.entity_code == "E001" for f in filtered.findings)
        assert len(filtered.findings) < 28

        export = svc.export_review(filtered, path=tmp_path / "filtered.csv", format="csv")
        assert export.success
        df = pd.read_csv(export.path)
        assert all(df["entity_code"] == "E001")

    def test_filtered_full_review(self, tmp_path):
        svc = MythosService()
        spec = build_filter_spec(entities=["E001", "E002"])
        reports = svc.full_review(
            FIXTURES / "sample_cy.xml",
            FIXTURES / "sample_py.xml",
            output_path=tmp_path / "filtered_full.xlsx",
            filter_spec=spec,
        )
        review = reports["_review_report"]
        assert all(f.entity_code in {"E001", "E002"} for f in review.findings)

    def test_filtered_by_category(self, tmp_path):
        svc = MythosService()
        spec = build_filter_spec(categories=["flow"])
        result = svc.review(FIXTURES / "sample_cy.xml", prior=FIXTURES / "sample_py.xml")
        filtered = filter_review_report(result.report, spec)
        assert all(f.category == "flow" for f in filtered.findings)


# ═══════════════════════════════════════════════════════════════════════════════
# REGRESSION — Data Integrity
# ═══════════════════════════════════════════════════════════════════════════════


class TestDataIntegrity:
    """Verify no data loss or corruption through the pipeline."""

    def test_entity_count_matches_parser(self):
        parser = EFileParser(FIXTURES / "sample_cy.xml")
        subs = parser.list_subsidiaries()
        engine = ReviewEngine()
        report = engine.review(FIXTURES / "sample_cy.xml", FIXTURES / "sample_py.xml")
        assert report.entity_count == len(subs)

    def test_findings_sorted_by_severity(self):
        engine = ReviewEngine()
        report = engine.review(FIXTURES / "sample_cy.xml", FIXTURES / "sample_py.xml")
        severities = [f.severity for f in report.findings]
        severity_order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
        for i in range(len(severities) - 1):
            assert severity_order[severities[i]] <= severity_order[severities[i + 1]]

    def test_all_checks_have_valid_check_id(self):
        engine = ReviewEngine()
        report = engine.review(FIXTURES / "sample_cy.xml", FIXTURES / "sample_py.xml")
        for f in report.findings:
            assert f.check_id.startswith(("FLO-", "CMP-", "RSN-", "ROL-", "XSC-", "XFM-"))

    def test_all_entities_have_names(self):
        engine = ReviewEngine()
        report = engine.review(FIXTURES / "sample_cy.xml", FIXTURES / "sample_py.xml")
        for f in report.findings:
            assert f.entity_name != ""
            assert f.entity_code != ""

    def test_report_summary_consistent(self):
        engine = ReviewEngine()
        report = engine.review(FIXTURES / "sample_cy.xml", FIXTURES / "sample_py.xml")
        report.compute_summary()
        s = report.summary
        assert s["total_findings"] == len(report.findings)
        assert s["by_severity"]["HIGH"] + s["by_severity"]["MEDIUM"] + s["by_severity"]["LOW"] == len(report.findings)
