"""Unit and integration tests for the Reconciler module.

Covers:
- Tolerance-based field comparison (_compare_field)
- Sign-flip handling (SIGN_FLIP_FIELDS)
- NaN / None / zero edge cases
- Two-way reconciliation (WB vs XML)
- Three-way reconciliation (OIT vs WB vs XML)
- Summary computation
- Report generation (ReconciliationReport, ThreeWayReport)
"""

import math
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from lab.xml_parser.reconciler import (
    ComparisonResult,
    FieldResult,
    Reconciler,
    ReconciliationReport,
    SIGN_FLIP_FIELDS,
    ThreeWayReport,
)


# ═══════════════════════════════════════════════════════════════════════════════
# UNIT TESTS — _compare_field()
# ═══════════════════════════════════════════════════════════════════════════════


class TestCompareField:
    """Direct tests for the tolerance comparison engine."""

    @pytest.fixture
    def rec(self):
        return Reconciler(tolerance=1.0)

    def test_both_none_passes(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "SomeField", None, None)
        assert r.status == "PASS"
        assert r.delta == 0.0

    def test_both_nan_passes(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "SomeField", float("nan"), float("nan"))
        assert r.status == "PASS"

    def test_exact_match(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", 1000.0, 1000.0)
        assert r.status == "PASS"
        assert r.delta == 0.0

    def test_within_tolerance(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", 1000.0, 1000.5)
        assert r.status == "PASS"
        assert r.delta == 0.5

    def test_at_tolerance_boundary(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", 1000.0, 1001.0)
        assert r.status == "PASS"
        assert r.delta == 1.0

    def test_exceeds_tolerance(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", 1000.0, 1002.0)
        assert r.status == "FAIL"
        assert r.delta == 2.0

    def test_negative_values(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", -500.0, -501.0)
        assert r.status == "PASS"
        assert r.delta == 1.0

    def test_string_exact_match(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", "USD", "USD")
        assert r.status == "PASS"

    def test_string_mismatch(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", "USD", "EUR")
        assert r.status == "TYPE_MISMATCH"

    def test_string_whitespace_stripped(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", " USD ", "USD")
        assert r.status == "PASS"

    def test_dormant_pattern_xml_nan_wb_zero(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", 0.0, float("nan"))
        assert r.status == "DORMANT_OK"

    def test_dormant_pattern_xml_none_wb_zero(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", 0, None)
        assert r.status == "DORMANT_OK"

    def test_xml_nan_wb_nonzero_fails(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", 1000.0, float("nan"))
        assert r.status == "FAIL"

    def test_large_tolerance(self):
        rec = Reconciler(tolerance=100.0)
        r = rec._compare_field("E001", "sch_h_fc", "Field", 1000.0, 1050.0)
        assert r.status == "PASS"
        assert r.delta == 50.0

    def test_zero_tolerance(self):
        rec = Reconciler(tolerance=0.0)
        r = rec._compare_field("E001", "sch_h_fc", "Field", 1000.0, 1000.01)
        assert r.status == "FAIL"

    def test_numeric_string_conversion(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", "1000", "1000")
        assert r.status == "PASS"

    def test_numeric_string_with_delta(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "Field", "1000", "1005")
        assert r.status == "FAIL"
        assert r.delta == 5.0


class TestSignFlipHandling:
    """Tests for the SIGN_FLIP_FIELDS special handling."""

    @pytest.fixture
    def rec(self):
        return Reconciler(tolerance=1.0)

    def test_sign_flip_field_passes_when_abs_match(self, rec):
        field_name = "TotalNetSubtractionsAmt"
        assert field_name in SIGN_FLIP_FIELDS
        r = rec._compare_field("E001", "sch_h_fc", field_name, 500.0, -500.0)
        assert r.status == "PASS_SIGN_FLIP"

    def test_sign_flip_field_fails_when_values_differ(self, rec):
        field_name = "TotalNetSubtractionsAmt"
        r = rec._compare_field("E001", "sch_h_fc", field_name, 500.0, -600.0)
        assert r.status == "FAIL"
        assert r.delta == 100.0 or r.delta == 1100.0

    def test_non_sign_flip_field_fails_on_sign_difference(self, rec):
        r = rec._compare_field("E001", "sch_h_fc", "RegularField", 500.0, -500.0)
        assert r.status == "FAIL"
        assert r.delta == 1000.0

    def test_all_sign_flip_fields_recognized(self):
        expected = {
            "TotalNetSubtractionsAmt",
            "OtherAdjustmentsNetSbtrctnAmt",
            "TotalTaxInFunctionalCurAmt",
            "TotalTaxInUSDollarsAmt",
            "SchEIncomeSubjectToTaxAmt",
        }
        assert SIGN_FLIP_FIELDS == expected


# ═══════════════════════════════════════════════════════════════════════════════
# UNIT TESTS — FieldResult and ReconciliationReport
# ═══════════════════════════════════════════════════════════════════════════════


class TestFieldResult:
    def test_default_status_is_pass(self):
        r = FieldResult("E001", "sch_h_fc", "Field", "1")
        assert r.status == "PASS"

    def test_all_fields_stored(self):
        r = FieldResult("E001", "sch_h_fc", "Income", "1a", 1000, 1001, 1.0, "PASS")
        assert r.entity_code == "E001"
        assert r.schedule == "sch_h_fc"
        assert r.field_name == "Income"
        assert r.line == "1a"
        assert r.wb_value == 1000
        assert r.xml_value == 1001
        assert r.delta == 1.0


class TestReconciliationReport:
    @pytest.fixture
    def sample_results(self):
        return [
            FieldResult("E001", "sch_h_fc", "F1", "1", 100, 100, 0, "PASS"),
            FieldResult("E001", "sch_h_fc", "F2", "2", 200, 300, 100, "FAIL"),
            FieldResult("E002", "sch_h_fc", "F1", "1", 0, None, None, "DORMANT_OK"),
            FieldResult("E003", "sch_h_fc", "F1", "1", 500, None, None, "XML_MISSING"),
        ]

    def test_to_dataframe(self, sample_results):
        report = ReconciliationReport(results=sample_results)
        df = report.to_dataframe()
        assert len(df) == 4
        assert "entity_code" in df.columns
        assert "status" in df.columns

    def test_failures_only(self, sample_results):
        report = ReconciliationReport(results=sample_results)
        fails = report.failures_only()
        assert len(fails) == 1
        assert fails.iloc[0]["entity_code"] == "E001"
        assert fails.iloc[0]["field_name"] == "F2"

    def test_empty_report(self):
        report = ReconciliationReport()
        assert report.to_dataframe().empty
        assert report.failures_only().empty

    def test_to_excel(self, sample_results, tmp_path):
        report = ReconciliationReport(results=sample_results)
        report.summary = {"total": 4, "pass_count": 1, "fail_count": 1}
        out = tmp_path / "test_recon.xlsx"
        report.to_excel(out)
        assert out.exists()
        assert out.stat().st_size > 0

        import openpyxl
        wb = openpyxl.load_workbook(out, read_only=True)
        assert "All Results" in wb.sheetnames
        assert "Failures" in wb.sheetnames
        assert "Summary" in wb.sheetnames


# ═══════════════════════════════════════════════════════════════════════════════
# UNIT TESTS — _compute_summary()
# ═══════════════════════════════════════════════════════════════════════════════


class TestComputeSummary:
    @pytest.fixture
    def rec(self):
        return Reconciler(tolerance=1.0)

    def test_empty_results(self, rec):
        s = rec._compute_summary([])
        assert s["total"] == 0
        assert s["pass_rate"] == 0.0

    def test_all_pass(self, rec):
        results = [
            FieldResult("E001", "s", "f", "", 100, 100, 0, "PASS"),
            FieldResult("E002", "s", "f", "", 200, 200, 0, "PASS"),
        ]
        s = rec._compute_summary(results)
        assert s["total"] == 2
        assert s["pass_count"] == 2
        assert s["fail_count"] == 0
        assert s["pass_rate"] == 1.0

    def test_mixed_results(self, rec):
        results = [
            FieldResult("E001", "s", "f1", "", 100, 100, 0, "PASS"),
            FieldResult("E001", "s", "f2", "", 200, 300, 100, "FAIL"),
            FieldResult("E002", "s", "f1", "", 0, None, None, "DORMANT_OK"),
            FieldResult("E003", "s", "f1", "", 500, None, None, "XML_MISSING"),
        ]
        s = rec._compute_summary(results)
        assert s["total"] == 4
        assert s["pass_count"] == 1
        assert s["fail_count"] == 1
        assert s["dormant_ok"] == 1
        assert s["xml_missing"] == 1
        assert s["pass_rate"] == pytest.approx(2 / 3, abs=0.01)

    def test_by_schedule_breakdown(self, rec):
        results = [
            FieldResult("E001", "sch_h_fc", "f1", "", 100, 100, 0, "PASS"),
            FieldResult("E001", "sch_h_fc", "f2", "", 200, 300, 100, "FAIL"),
            FieldResult("E001", "sch_f_usd", "f1", "", 400, 400, 0, "PASS"),
        ]
        s = rec._compute_summary(results)
        assert "by_schedule" in s
        assert s["by_schedule"]["sch_h_fc"]["pass"] == 1
        assert s["by_schedule"]["sch_h_fc"]["fail"] == 1
        assert s["by_schedule"]["sch_f_usd"]["pass"] == 1


# ═══════════════════════════════════════════════════════════════════════════════
# INTEGRATION TESTS — Two-Way Reconciliation
# ═══════════════════════════════════════════════════════════════════════════════


class TestTwoWayReconciliation:
    """Test the full reconcile() flow with mock data."""

    @pytest.fixture
    def rec(self):
        return Reconciler(tolerance=1.0)

    @pytest.fixture
    def mock_wb_data(self):
        """Minimal WorkbookData-like object with schedules dict."""

        class MockWBData:
            def __init__(self):
                self.schedules = {
                    "sch_h_fc": pd.DataFrame(
                        {
                            "ForeignCYNetIncomePerBooksAmt": [1000000, -500000, 0],
                            "CurrentEarningsAndProfitsAmt": [800000, -400000, 0],
                            "ExchangeRt": [0.85, 0.91, 1.0],
                        },
                        index=["E001", "E002", "E003"],
                    ),
                }

        return MockWBData()

    @pytest.fixture
    def mock_parser(self):
        """Mock parser that returns a pre-built DataFrame for extract_form."""

        class MockParser:
            def extract_form(self, form_name):
                if form_name == "IRS5471ScheduleH":
                    return pd.DataFrame(
                        {
                            "_reference_id": ["E001", "E002", "E003"],
                            "IRS5471ScheduleH_ForeignCYNetIncomePerBooksAmt": [
                                1000000, -500000, 0
                            ],
                            "IRS5471ScheduleH_CurrentEarningsAndProfitsAmt": [
                                800001, -400000, 0  # E001 is 1 off
                            ],
                            "IRS5471ScheduleH_ExchangeRt": [0.85, 0.91, 1.0],
                        }
                    )
                return pd.DataFrame()

            def to_dataframe(self):
                return pd.DataFrame()

        return MockParser()

    def test_reconcile_basic(self, rec, mock_wb_data, mock_parser):
        report = rec.reconcile(mock_wb_data, mock_parser, schedules=["sch_h_fc"])
        assert isinstance(report, ReconciliationReport)
        assert len(report.results) > 0

    def test_reconcile_pass_within_tolerance(self, rec, mock_wb_data, mock_parser):
        report = rec.reconcile(mock_wb_data, mock_parser, schedules=["sch_h_fc"])
        ep_results = [r for r in report.results if "Earnings" in r.field_name and r.entity_code == "E001"]
        if ep_results:
            assert ep_results[0].status == "PASS"
            assert ep_results[0].delta == 1.0

    def test_reconcile_summary_populated(self, rec, mock_wb_data, mock_parser):
        report = rec.reconcile(mock_wb_data, mock_parser, schedules=["sch_h_fc"])
        assert "total" in report.summary
        assert "pass_count" in report.summary
        assert "fail_count" in report.summary
        assert "pass_rate" in report.summary

    def test_reconcile_no_matching_schedules(self, rec, mock_wb_data, mock_parser):
        report = rec.reconcile(mock_wb_data, mock_parser, schedules=["nonexistent"])
        assert len(report.results) == 0
        assert report.summary["total"] == 0


# ═══════════════════════════════════════════════════════════════════════════════
# UNIT TESTS — Three-Way Reconciliation
# ═══════════════════════════════════════════════════════════════════════════════


class TestThreeWayComparison:
    """Test _compare_three_values() for all status paths."""

    @pytest.fixture
    def rec(self):
        return Reconciler(tolerance=1.0)

    def test_all_none_passes(self, rec):
        r = rec._compare_three_values("E001", "s", "f", None, None, None)
        assert r.status == "PASS"

    def test_all_match(self, rec):
        r = rec._compare_three_values("E001", "s", "f", 1000, 1000, 1000)
        assert r.status == "PASS"

    def test_all_within_tolerance(self, rec):
        r = rec._compare_three_values("E001", "s", "f", 1000, 1000.5, 1000.8)
        assert r.status == "PASS"

    def test_fail_oit_wb(self, rec):
        r = rec._compare_three_values("E001", "s", "f", 1000, 1100, 1000)
        assert r.status == "FAIL_OIT_WB"

    def test_fail_oit_xml(self, rec):
        r = rec._compare_three_values("E001", "s", "f", 1000, 1000, 1100)
        assert r.status == "FAIL_OIT_XML"

    def test_fail_wb_xml(self, rec):
        # OIT=1000, WB=999, XML=1001 → OIT/WB=1(pass), OIT/XML=1(pass), WB/XML=2(fail)
        r = rec._compare_three_values("E001", "s", "f", 1000, 999, 1001)
        assert r.status == "FAIL_WB_XML"

    def test_fail_all(self, rec):
        r = rec._compare_three_values("E001", "s", "f", 1000, 2000, 3000)
        assert r.status == "FAIL_ALL"

    def test_oit_missing(self, rec):
        r = rec._compare_three_values("E001", "s", "f", None, 1000, 1000)
        assert r.status == "OIT_MISSING"

    def test_wb_missing(self, rec):
        r = rec._compare_three_values("E001", "s", "f", 1000, None, 1000)
        assert r.status == "WB_MISSING"

    def test_xml_missing(self, rec):
        r = rec._compare_three_values("E001", "s", "f", 1000, 1000, None)
        assert r.status == "XML_MISSING"

    def test_dormant_ok_when_present_values_zero(self, rec):
        r = rec._compare_three_values("E001", "s", "f", None, 0, 0)
        assert r.status == "DORMANT_OK"

    def test_type_mismatch(self, rec):
        r = rec._compare_three_values("E001", "s", "f", "abc", "def", "ghi")
        assert r.status == "TYPE_MISMATCH"

    def test_sign_flip_passes(self, rec):
        r = rec._compare_three_values(
            "E001", "s", "TotalNetSubtractionsAmt", 500, 500, -500
        )
        assert r.status == "PASS_SIGN_FLIP"


class TestThreeWayReport:
    @pytest.fixture
    def sample_results(self):
        return [
            ComparisonResult("E001", "s", "f1", "1", 100, 100, 100, 0, 0, 0, "PASS"),
            ComparisonResult("E001", "s", "f2", "2", 100, 200, 300, 100, 200, 100, "FAIL_ALL"),
            ComparisonResult("E002", "s", "f1", "1", None, 0, 0, None, None, 0, "DORMANT_OK"),
        ]

    def test_to_dataframe(self, sample_results):
        report = ThreeWayReport(results=sample_results)
        df = report.to_dataframe()
        assert len(df) == 3
        assert "oit_value" in df.columns
        assert "wb_value" in df.columns
        assert "xml_value" in df.columns

    def test_failures_only(self, sample_results):
        report = ThreeWayReport(results=sample_results)
        fails = report.failures_only()
        assert len(fails) == 1
        assert fails.iloc[0]["status"] == "FAIL_ALL"

    def test_by_schedule(self, sample_results):
        report = ThreeWayReport(results=sample_results)
        bs = report.by_schedule()
        assert not bs.empty
        assert "PASS" in bs.columns

    def test_by_entity(self, sample_results):
        report = ThreeWayReport(results=sample_results)
        be = report.by_entity()
        assert "E001" in be.index

    def test_empty_report(self):
        report = ThreeWayReport()
        assert report.to_dataframe().empty
        assert report.failures_only().empty

    def test_to_excel(self, sample_results, tmp_path):
        report = ThreeWayReport(results=sample_results)
        report.summary = {"total": 3, "pass_count": 1, "fail_count": 1}
        out = tmp_path / "three_way.xlsx"
        report.to_excel(out)
        assert out.exists()

        import openpyxl
        wb = openpyxl.load_workbook(out, read_only=True)
        assert "Three-Way Results" in wb.sheetnames
        assert "Summary" in wb.sheetnames


class TestThreeWaySummary:
    @pytest.fixture
    def rec(self):
        return Reconciler(tolerance=1.0)

    def test_empty(self, rec):
        s = rec._compute_three_way_summary([])
        assert s["total"] == 0

    def test_mixed_statuses(self, rec):
        results = [
            ComparisonResult("E001", "s", "f1", "", status="PASS"),
            ComparisonResult("E001", "s", "f2", "", status="PASS_SIGN_FLIP"),
            ComparisonResult("E001", "s", "f3", "", status="FAIL_ALL"),
            ComparisonResult("E001", "s", "f4", "", status="FAIL_OIT_WB"),
            ComparisonResult("E001", "s", "f5", "", status="FAIL_OIT_XML"),
            ComparisonResult("E001", "s", "f6", "", status="FAIL_WB_XML"),
            ComparisonResult("E001", "s", "f7", "", status="DORMANT_OK"),
            ComparisonResult("E001", "s", "f8", "", status="OIT_MISSING"),
            ComparisonResult("E001", "s", "f9", "", status="WB_MISSING"),
            ComparisonResult("E001", "s", "f10", "", status="XML_MISSING"),
        ]
        s = rec._compute_three_way_summary(results)
        assert s["total"] == 10
        assert s["pass_count"] == 2
        assert s["dormant_ok"] == 1
        assert s["fail_all"] == 1
        assert s["fail_oit_wb"] == 1
        assert s["fail_oit_xml"] == 1
        assert s["fail_wb_xml"] == 1
        assert s["oit_missing"] == 1
        assert s["wb_missing"] == 1
        assert s["xml_missing"] == 1
        assert s["fail_count"] == 4
        comparable = 10 - 3  # exclude 3 missing
        expected_rate = (2 + 1) / comparable
        assert s["pass_rate"] == pytest.approx(expected_rate, abs=0.01)
